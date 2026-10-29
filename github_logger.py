import base64
import io
from datetime import datetime
from urllib.parse import quote

import requests
import streamlit as st
from openpyxl import Workbook, load_workbook


HEADERS = [
    "Timestamp",
    "User Name",
    "Department",
    "Conversation ID",
    "Question",
    "Answer",
    "Answer Type",
    "Sources",
    "Public IP",
    "Approximate Location",
]


def log_interaction(
    user_name,
    department,
    conversation_id,
    question,
    answer,
    answer_type,
    sources=None,
    public_ip="",
    approximate_location="",
):
    """Append an interaction to the GitHub-hosted Excel workbook."""

    token = st.secrets["GITHUB_TOKEN"]
    repo = st.secrets["GITHUB_REPO"]
    branch = st.secrets.get("GITHUB_BRANCH", "main")
    file_path = st.secrets.get(
        "GITHUB_XLSX_PATH",
        "data/bidbuddy_conversations.xlsx",
    )

    api_url = (
        f"https://api.github.com/repos/{repo}/contents/"
        f"{quote(file_path, safe='/')}"
    )

    headers = {
        "Authorization": f"Bearer {token}",
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
    }

    response = requests.get(
        api_url,
        headers=headers,
        params={"ref": branch},
        timeout=30,
    )

    file_sha = None

    if response.status_code == 200:
        file_info = response.json()
        file_sha = file_info["sha"]

        file_bytes = base64.b64decode(file_info["content"])
        workbook = load_workbook(io.BytesIO(file_bytes))
        sheet = workbook.active

        # Add any missing columns to the existing header row.
        existing_headers = {
            str(sheet.cell(row=1, column=col).value).strip(): col
            for col in range(1, sheet.max_column + 1)
            if sheet.cell(row=1, column=col).value is not None
        }

        for header in HEADERS:
            if header not in existing_headers:
                sheet.cell(
                    row=1,
                    column=sheet.max_column + 1,
                    value=header,
                )

        # Map header names to their actual column positions.
        column_map = {
            str(sheet.cell(row=1, column=col).value).strip(): col
            for col in range(1, sheet.max_column + 1)
            if sheet.cell(row=1, column=col).value is not None
        }

    elif response.status_code == 404:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Conversations"
        sheet.append(HEADERS)

        column_map = {
            header: index + 1
            for index, header in enumerate(HEADERS)
        }

    else:
        raise RuntimeError(
            f"Could not read GitHub workbook (HTTP "
            f"{response.status_code}). Check repository settings "
            "and token permissions."
        )

    source_text = (
        "\n".join(sources)
        if isinstance(sources, (list, tuple))
        else (sources or "")
    )

    row_data = {
        "Timestamp": datetime.now().astimezone().isoformat(
            timespec="seconds"
        ),
        "User Name": user_name,
        "Department": department,
        "Conversation ID": conversation_id,
        "Question": question,
        "Answer": answer,
        "Answer Type": answer_type,
        "Sources": source_text,
        "Public IP": public_ip or "",
        "Approximate Location": approximate_location or "",
    }

    # Write by header name, so column order remains safe.
    new_row = sheet.max_row + 1

    for header, value in row_data.items():
        sheet.cell(
            row=new_row,
            column=column_map[header],
            value=value,
        )

    buffer = io.BytesIO()
    workbook.save(buffer)

    encoded_content = base64.b64encode(
        buffer.getvalue()
    ).decode("utf-8")

    payload = {
        "message": "Log BidBuddy conversation",
        "content": encoded_content,
        "branch": branch,
    }

    if file_sha:
        payload["sha"] = file_sha

    save_response = requests.put(
        api_url,
        headers=headers,
        json=payload,
        timeout=30,
    )

    if save_response.status_code not in (200, 201):
        raise RuntimeError(
            f"Could not save GitHub workbook (HTTP "
            f"{save_response.status_code}): "
            f"{save_response.text[:500]}"
        )

    return True
